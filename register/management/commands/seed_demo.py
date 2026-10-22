import secrets
from decimal import Decimal
from pathlib import Path

from django.conf import settings
from django.contrib.auth import get_user_model
from django.core.management.base import BaseCommand
from openpyxl import Workbook

from register.models import Afdeling, Borger

AFDELINGER = [("Snedkeriet", "Rasmus"), ("Montagen", "Pia"), ("Trykkeriet", "Bent"), ("Café", "Ulla")]

# navn, timer:minutter, afdeling, km, transport, zoner, sats
PERSONER = [
    ("Person 1", "79:25", "Snedkeriet", "4.7", "selv", 2, "30"),
    ("Person 2", "100:58", "Snedkeriet", "2.0", "taxa", None, "20"),
    ("Person 3", "76:07", "Montagen", "13.2", "selv", None, "20"),
    ("Person 4", "99:11", "Montagen", "3.2", "selv", 2, "20"),
    ("Person 5", "88:27", "Trykkeriet", "1.9", "ingen", None, "15"),
    ("Person 6", "8:49", "Café", "12.8", "selv", 3, "15"),
    ("Person 7", "49:06", "Café", "6.0", "selv", 2, "15"),
    ("Person 8", "35:21", "Trykkeriet", "0.9", "ingen", None, "12"),
    ("Person 9", "50:08", "Snedkeriet", "2.1", "selv", 2, "12"),
    ("Person 10", "56:06", "Montagen", "8.4", "taxa", None, "20"),
    ("Person 11", "7:58", "Café", "1.8", "selv", 2, "15"),
    ("Person 12", "81:35", "Snedkeriet", "13.2", "selv", 3, "20"),
]


class Command(BaseCommand):
    help = "Opretter fiktive demodata (ingen rigtige borgere), en demobruger og en eksempelfil i stil med Punchly-eksporten."

    def add_arguments(self, parser):
        parser.add_argument("--bruger", default="demo", help="Brugernavn til demobrugeren")

    def handle(self, *args, **opts):
        afd = {n: Afdeling.objects.get_or_create(navn=n, defaults={"lonseddel_modtager": m})[0] for n, m in AFDELINGER}
        for i, (navn, _, a, km, transport, zoner, sats) in enumerate(PERSONER, start=1):
            Borger.objects.get_or_create(
                navn=navn,
                defaults=dict(
                    tj_nr=f"9{i:03d}",
                    cpr=f"000000-{i:04d}",
                    adresse=f"Eksempelvej {i}, 9800 Hjørring",
                    afdeling=afd[a],
                    arbejdssted="Eksempelværkstedet",
                    afstand_km=Decimal(km),
                    transport=transport,
                    zoner=zoner,
                    sats=Decimal(sats),
                ),
            )

        mappe = Path(settings.BASE_DIR) / "eksempel"
        mappe.mkdir(exist_ok=True)
        wb = Workbook()
        ws = wb.active
        ws.title = "Ark1"
        ws["C4"], ws["D4"] = "Navn", "Total timer (timer:minutter)"
        for r, (navn, tid, *_rest) in enumerate(PERSONER, start=5):
            ws.cell(row=r, column=3, value=navn)
            ws.cell(row=r, column=4, value=tid)
        sidste = 5 + len(PERSONER)
        ws.cell(row=sidste, column=3, value="Person 13")
        ws.cell(row=sidste, column=4, value="56:06")
        wb.save(mappe / "punchly-eksempel.xlsx")
        self.stdout.write(self.style.SUCCESS(f"Eksempelfil gemt: {mappe / 'punchly-eksempel.xlsx'}"))

        User = get_user_model()
        if not User.objects.filter(username=opts["bruger"]).exists():
            adgangskode = secrets.token_urlsafe(12)
            User.objects.create_user(opts["bruger"], password=adgangskode, is_staff=True, is_superuser=True)
            self.stdout.write(self.style.SUCCESS(f"Bruger '{opts['bruger']}' oprettet med adgangskode: {adgangskode}"))
        else:
            self.stdout.write(f"Bruger '{opts['bruger']}' findes allerede.")

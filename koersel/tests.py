import hashlib
import re
import shutil
import tempfile
from datetime import date, time, timedelta
from decimal import Decimal
from html.parser import HTMLParser
from io import BytesIO
from unittest import mock

from django.conf import settings
from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from django.urls import reverse
from django.utils import timezone
from openpyxl import Workbook
from reportlab import rl_config

from audit.models import AuditLog
from beregning.pdf import LAYOUT_VERSION, lav_rapport
from register.models import Borger

from . import services
from .importer import ImportFejl, laes_timefil
from .models import Koersel, Linje


def xlsx(raekker, header=("Navn", "Total timer (timer:minutter)"), navn="timer.xlsx"):
    """Rækker er (navn, timer) eller (navn, timer, dage); header med tre felter giver en dagskolonne."""
    wb = Workbook()
    ws = wb.active
    for j, h in enumerate(header):
        ws.cell(row=4, column=3 + j, value=h)
    for i, raekke in enumerate(raekker, start=5):
        for j, v in enumerate(raekke):
            ws.cell(row=i, column=3 + j, value=v)
    buf = BytesIO()
    wb.save(buf)
    return SimpleUploadedFile(navn, buf.getvalue())


class _SynligTekst(HTMLParser):
    """Samler det en bruger kan læse i en side: tekst og de attributter, der vises eller læses op."""

    SYNLIGE = ("aria-label", "placeholder", "title", "alt")

    def __init__(self):
        super().__init__()
        self.dele = []
        self._skjult = 0

    def handle_starttag(self, tag, attrs):
        if tag in ("script", "style"):
            self._skjult += 1
        self.dele.extend(v for k, v in attrs if k in self.SYNLIGE and v)

    def handle_endtag(self, tag):
        if tag in ("script", "style") and self._skjult:
            self._skjult -= 1

    def handle_data(self, data):
        if not self._skjult:
            self.dele.append(data)


def synlig_tekst(html: str) -> str:
    p = _SynligTekst()
    p.feed(html)
    return " ".join(p.dele)


class ImporterTests(TestCase):
    def test_laeser_tekst_tider(self):
        r = laes_timefil(xlsx([("Person 1", "79:25"), ("Person 2", "8:49")]))
        self.assertEqual([(x.navn, x.minutter) for x in r], [("Person 1", 79 * 60 + 25), ("Person 2", 8 * 60 + 49)])

    def test_laeser_excel_tidsceller(self):
        r = laes_timefil(xlsx([("A", time(8, 30)), ("B", timedelta(hours=100, minutes=58)), ("C", 0.5)]))
        self.assertEqual([x.minutter for x in r], [510, 6058, 720])

    def test_afviser_forkert_filtype(self):
        with self.assertRaisesMessage(ImportFejl, ".xlsx"):
            laes_timefil(SimpleUploadedFile("timer.csv", b"Navn;Timer"))

    def test_afviser_ikke_excel_indhold(self):
        with self.assertRaises(ImportFejl):
            laes_timefil(SimpleUploadedFile("timer.xlsx", b"ikke en excel-fil"))

    def test_afviser_beskadiget_excel_fil_uden_at_crashe(self):
        god = bytearray(xlsx([("A", "1:00")]).read())
        for i in range(120, 260):
            god[i] = 0
        with self.assertRaisesMessage(ImportFejl, "kan ikke læses"):
            laes_timefil(SimpleUploadedFile("timer.xlsx", bytes(god)))

    def test_mangler_kolonner(self):
        with self.assertRaisesMessage(ImportFejl, "Kunne ikke finde kolonnerne"):
            laes_timefil(xlsx([("A", "1:00")], header=("Person", "Antal")))

    def test_ugyldigt_tidsformat_peger_paa_raekken(self):
        with self.assertRaisesMessage(ImportFejl, "Række 6 (B)"):
            laes_timefil(xlsx([("A", "1:00"), ("B", "meget")]))

    def test_laeser_antal_arbejdsdage(self):
        r = laes_timefil(xlsx([("A", "1:00", 17), ("B", "2:00", "5"), ("C", "3:00", 31.0)], header=("Navn", "Total timer", "Antal arbejdsdage")))
        self.assertEqual([x.dage for x in r], [17, 5, 31])

    def test_laeser_ogsaa_antal_dage(self):
        r = laes_timefil(xlsx([("A", "1:00", 17), ("B", "2:00", "5")], header=("Navn", "Total timer", "Antal dage")))
        self.assertEqual([x.dage for x in r], [17, 5])

    def test_dagskolonnen_ses_bort_fra_store_bogstaver_og_mellemrum(self):
        for overskrift in ("  ANTAL arbejdsdage ", "antal ARBEJDSDAGE", " Antal   Arbejdsdage", "  ANTAL dage ", "antal DAGE"):
            with self.subTest(overskrift=overskrift):
                r = laes_timefil(xlsx([("A", "1:00", 4)], header=("Navn", "Total timer", overskrift)))
                self.assertEqual(r[0].dage, 4)

    def test_dagskolonnen_kan_staa_foer_navn_og_timer(self):
        r = laes_timefil(xlsx([(8, "A", "1:00")], header=("Antal arbejdsdage", "Navn", "Total timer (timer:minutter)")))
        self.assertEqual([(x.navn, x.minutter, x.dage) for x in r], [("A", 60, 8)])

    def test_tom_celle_giver_ingen_dage(self):
        r = laes_timefil(xlsx([("A", "1:00", None), ("B", "2:00", "  "), ("C", "3:00", 9)], header=("Navn", "Total timer", "Antal dage")))
        self.assertEqual([x.dage for x in r], [None, None, 9])

    def test_manglende_dagskolonne_giver_ingen_dage_og_ingen_fejl(self):
        r = laes_timefil(xlsx([("A", "1:00"), ("B", "2:00")]))
        self.assertEqual([(x.minutter, x.dage) for x in r], [(60, None), (120, None)])

    def test_ugyldige_dage_afviser_hele_filen_med_raekke_og_vaerdi(self):
        for overskrift in ("Antal arbejdsdage", "Antal dage"):
            for ugyldig in (0, 32, -1, 2.5, "mange", "4,5", True):
                with self.subTest(overskrift=overskrift, ugyldig=ugyldig):
                    header = ("Navn", "Total timer", overskrift)
                    with self.assertRaisesMessage(ImportFejl, f"Række 6 (B): '{ugyldig}' i kolonnen '{overskrift}' er ikke et gyldigt antal dage"):
                        laes_timefil(xlsx([("A", "1:00", 3), ("B", "2:00", ugyldig)], header=header))


TEST_MEDIA = tempfile.mkdtemp(prefix="borgerlon-test-")


@override_settings(MEDIA_ROOT=TEST_MEDIA)
class FlowTests(TestCase):
    @classmethod
    def tearDownClass(cls):
        super().tearDownClass()
        shutil.rmtree(TEST_MEDIA, ignore_errors=True)

    def setUp(self):
        self.bruger = get_user_model().objects.create_user("tester", password="x-test-adgangskode-1")
        self.client.force_login(self.bruger)
        self.p1 = Borger.objects.create(navn="Person 1", tj_nr="9001", cpr="000000-0001", sats=Decimal("20"), afstand_km=Decimal("4.7"), transport="selv", zoner=2)
        self.p2 = Borger.objects.create(navn="Person 2", tj_nr="9002", cpr="000000-0002", sats=Decimal("20"), afstand_km=Decimal("2"), transport="taxa")
        self.p3 = Borger.objects.create(navn="Person 3", tj_nr="9003", cpr="000000-0003", sats=Decimal("20"), afstand_km=Decimal("13.2"), transport="selv")

    def opret_og_upload(self, raekker):
        self.client.post(reverse("dashboard"), {"maaned": "2026-10"})
        k = Koersel.objects.get()
        svar = self.client.post(reverse("upload", args=[k.pk]), {"fil": xlsx(raekker, navn="punchly.xlsx")})
        self.assertRedirects(svar, reverse("gennemse", args=[k.pk]))
        return k

    def undtagelser(self, k, **valg):
        data = {"handling": valg.pop("handling", "beregn")}
        for l in k.linjer.all():
            data.update({f"t_{l.pk}": l.transport, f"z_{l.pk}": l.zoner or ""})
            if l.individuel:
                data[f"i_{l.pk}"] = "on"
                data[f"d_{l.pk}"] = l.dage or ""
        data.update(valg)
        return self.client.post(reverse("undtagelser", args=[k.pk]), data)

    def test_upload_matcher_paa_navn(self):
        k = self.opret_og_upload([("Person 1", "44:12"), ("  person 2 ", "10:00"), ("Ukendt", "5:00")])
        status = {l.fil_navn.strip(): l.match_status for l in k.linjer.all()}
        self.assertEqual(status, {"Person 1": "matchet", "person 2": "matchet", "Ukendt": "ingen"})

    def test_upload_gemmer_hvem_og_hvornaar_og_viser_det(self):
        k = self.opret_og_upload([("Person 1", "44:12")])
        k.refresh_from_db()
        self.assertEqual(k.uploadet_af, self.bruger)
        self.assertIsNotNone(k.uploadet)
        svar = self.client.get(reverse("upload", args=[k.pk]))
        self.assertContains(svar, "Sidst uploadet: punchly.xlsx, af tester, ")
        self.assertContains(svar, timezone.localtime(k.uploadet).strftime("%d.%m.%Y kl. %H:%M"))

    def test_ny_upload_erstatter_fil_bruger_og_tidspunkt(self):
        k = self.opret_og_upload([("Person 1", "44:12")])
        foer = timezone.now() - timedelta(days=3)
        Koersel.objects.filter(pk=k.pk).update(uploadet=foer)
        anden = get_user_model().objects.create_user("anden", password="x-test-adgangskode-2")
        self.client.force_login(anden)
        self.client.post(reverse("upload", args=[k.pk]), {"fil": xlsx([("Person 2", "1:00")], navn="ny.xlsx")})
        k.refresh_from_db()
        self.assertEqual((k.filnavn, k.uploadet_af), ("ny.xlsx", anden))
        self.assertGreater(k.uploadet, foer)
        svar = self.client.get(reverse("upload", args=[k.pk]))
        self.assertContains(svar, "Sidst uploadet: ny.xlsx, af anden, ")
        self.assertNotContains(svar, "punchly.xlsx")

    def test_afvist_upload_aendrer_ikke_sidst_uploadet(self):
        k = self.opret_og_upload([("Person 1", "44:12")])
        foer = Koersel.objects.get(pk=k.pk)
        self.client.post(reverse("upload", args=[k.pk]), {"fil": SimpleUploadedFile("daarlig.xlsx", b"ikke en excel-fil")})
        efter = Koersel.objects.get(pk=k.pk)
        self.assertEqual((efter.filnavn, efter.uploadet_af, efter.uploadet), (foer.filnavn, foer.uploadet_af, foer.uploadet))

    def test_koersel_uploadet_foer_felterne_viser_kun_filnavn(self):
        k = Koersel.objects.create(maaned=date(2026, 9, 1), filnavn="gammel.xlsx")
        svar = self.client.get(reverse("upload", args=[k.pk]))
        self.assertContains(svar, "Sidst uploadet: gammel.xlsx<")
        self.assertNotContains(svar, ", af ")
        self.assertNotContains(svar, " kl. ")

    def test_slettet_bruger_viser_filnavn_og_tidspunkt_uden_navn(self):
        k = self.opret_og_upload([("Person 1", "44:12")])
        self.bruger.delete()
        self.client.force_login(get_user_model().objects.create_user("ny", password="x-test-adgangskode-3"))
        k.refresh_from_db()
        self.assertIsNone(k.uploadet_af)
        svar = self.client.get(reverse("upload", args=[k.pk]))
        self.assertContains(svar, "Sidst uploadet: punchly.xlsx, ")
        self.assertNotContains(svar, ", af ")
        self.assertContains(svar, " kl. ")

    def test_uden_upload_staar_der_ingen_fil_valgt(self):
        k = Koersel.objects.create(maaned=date(2026, 9, 1))
        self.assertContains(self.client.get(reverse("upload", args=[k.pk])), "Ingen fil valgt")

    def upload_med_dage(self, raekker):
        self.client.post(reverse("dashboard"), {"maaned": "2026-10"})
        k = Koersel.objects.get()
        fil = xlsx(raekker, header=("Navn", "Total timer", "Antal arbejdsdage"), navn="punchly.xlsx")
        self.assertRedirects(self.client.post(reverse("upload", args=[k.pk]), {"fil": fil}), reverse("gennemse", args=[k.pk]))
        return k

    def test_upload_gemmer_dage_fra_filen_men_bruger_dem_ikke_endnu(self):
        k = self.upload_med_dage([("Person 1", "44:12", 17), ("Person 2", "10:00", None)])
        self.assertEqual({l.fil_navn: (l.fil_dage, l.dage) for l in k.linjer.all()}, {"Person 1": (17, None), "Person 2": (None, None)})

    def test_upload_beskriver_antal_dage_som_valgfri_kolonne(self):
        k = Koersel.objects.create(maaned=date(2026, 9, 1))
        self.assertContains(self.client.get(reverse("upload", args=[k.pk])), 'Kolonnen "Antal arbejdsdage"')

    def test_undtagelser_forudfylder_dage_kun_ved_individuel_beregning(self):
        k = self.upload_med_dage([("Person 1", "44:12", 17), ("Person 3", "10:00", 8)])
        svar = self.client.get(reverse("undtagelser", args=[k.pk]))
        l1 = k.linjer.get(fil_navn="Person 1")
        self.assertContains(svar, f'name="d_{l1.pk}" data-felt="d" data-fil-dage="17" placeholder="Dage" aria-label="Fremmødedage for Person 1" value=""')
        l1.individuel = True
        l1.save()
        svar = self.client.get(reverse("undtagelser", args=[k.pk]))
        self.assertContains(svar, f'name="d_{l1.pk}" data-felt="d" data-fil-dage="17" placeholder="Dage" aria-label="Fremmødedage for Person 1" value="17"')

    def test_undtagelser_uden_dage_i_filen_er_feltet_tomt(self):
        k = self.upload_med_dage([("Person 1", "44:12", None)])
        l = k.linjer.get()
        l.individuel = True
        l.save()
        svar = self.client.get(reverse("undtagelser", args=[k.pk]))
        self.assertContains(svar, 'data-fil-dage="" placeholder="Dage" aria-label="Fremmødedage for Person 1" value=""')

    def test_rettede_dage_vinder_over_filens(self):
        k = self.upload_med_dage([("Person 1", "44:12", 17)])
        l = k.linjer.get()
        self.undtagelser(k, handling="gem", **{f"i_{l.pk}": "on", f"d_{l.pk}": "12"})
        svar = self.client.get(reverse("undtagelser", args=[k.pk]))
        self.assertContains(svar, 'data-fil-dage="17" placeholder="Dage" aria-label="Fremmødedage for Person 1" value="12"')

    def test_filens_dage_overlever_gem_uden_individuel_og_manuel_afklaring(self):
        k = self.upload_med_dage([("Person 1", "44:12", 17), ("Ukendt", "5:00", 9)])
        ukendt = k.linjer.get(fil_navn="Ukendt")
        self.client.post(reverse("gennemse", args=[k.pk]), {"handling": "afklar", "linje": ukendt.pk, "borger": self.p1.pk})
        ukendt.refresh_from_db()
        self.assertEqual((ukendt.borger, ukendt.fil_dage, ukendt.dage), (self.p1, 9, None))
        self.undtagelser(k, handling="gem", **{f"i_{ukendt.pk}": "on", f"d_{ukendt.pk}": "5"})
        ukendt.refresh_from_db()
        self.assertEqual((ukendt.individuel, ukendt.dage, ukendt.fil_dage), (True, 5, 9))
        data = {"handling": "gem", **{f"t_{l.pk}": l.transport for l in k.linjer.all()}, **{f"z_{l.pk}": l.zoner or "" for l in k.linjer.all()}}
        self.client.post(reverse("undtagelser", args=[k.pk]), data)
        ukendt.refresh_from_db()
        self.assertEqual((ukendt.individuel, ukendt.dage, ukendt.fil_dage), (False, None, 9))
        self.assertEqual(k.linjer.get(fil_navn="Person 1").fil_dage, 17)

    def test_beregning_bruger_dage_kun_ved_individuel_beregning(self):
        Borger.objects.create(navn="Person 4", tj_nr="9004", cpr="000000-0004", sats=Decimal("20"), afstand_km=Decimal("3.2"), transport="selv", zoner=2)
        k = self.upload_med_dage([("Person 1", "44:12", 17), ("Person 4", "10:00", 20)])
        self.assertEqual(self.undtagelser(k, handling="gem").status_code, 302)
        for l in k.linjer.all():
            self.assertIsNone(l.dage)
        p1, p4 = k.linjer.get(fil_navn="Person 1"), k.linjer.get(fil_navn="Person 4")
        self.undtagelser(k, handling="beregn", **{f"i_{p1.pk}": "on", f"d_{p1.pk}": "17", f"d_{p4.pk}": "20"})
        linjer = {x["navn"]: x["input"] for x in k.aktiv_beregning.data["linjer"]}
        self.assertEqual((linjer["Person 1"]["individuel"], linjer["Person 1"]["dage"]), (True, 17))
        self.assertEqual((linjer["Person 4"]["individuel"], linjer["Person 4"]["dage"]), (False, None))

    def test_flertydigt_navn_afklares_aldrig_automatisk(self):
        Borger.objects.create(navn="Person 1", tj_nr="9101")
        k = self.opret_og_upload([("Person 1", "1:00")])
        l = k.linjer.get()
        self.assertEqual(l.match_status, Linje.Match.FLERTYDIG)
        self.assertIsNone(l.borger)

    def test_uafklarede_raekker_blokerer_videre_flow(self):
        k = self.opret_og_upload([("Ukendt", "5:00")])
        svar = self.client.get(reverse("undtagelser", args=[k.pk]))
        self.assertRedirects(svar, reverse("gennemse", args=[k.pk]))
        svar = self.client.post(reverse("gennemse", args=[k.pk]), {"handling": "fortsaet"})
        self.assertRedirects(svar, reverse("gennemse", args=[k.pk]))

    def test_afklar_raekke_og_gem_rettelser_i_registret(self):
        k = self.opret_og_upload([("Ukendt", "5:00")])
        l = k.linjer.get()
        self.client.post(reverse("gennemse", args=[k.pk]), {"handling": "afklar", "linje": l.pk, "borger": self.p2.pk})
        l.refresh_from_db()
        self.assertEqual(l.borger, self.p2)
        self.client.post(reverse("gennemse", args=[k.pk]), {"handling": "gem", f"afstand_{l.pk}": "7,5", f"transport_{l.pk}": "ingen"})
        self.p2.refresh_from_db()
        self.assertEqual((self.p2.afstand_km, self.p2.transport), (Decimal("7.5"), "ingen"))
        self.assertTrue(AuditLog.objects.filter(handling="Ændrede medarbejder fra gennemsyn").exists())

    def test_afklar_uden_valgt_medarbejder_giver_besked_og_aendrer_intet(self):
        k = self.opret_og_upload([("Ukendt", "5:00")])
        l = k.linjer.get()
        antal_logs = AuditLog.objects.count()
        for data in ({"handling": "afklar", "linje": l.pk, "borger": ""}, {"handling": "afklar", "linje": l.pk}, {"handling": "afklar", "linje": l.pk, "borger": "abc"}):
            with self.subTest(data=data):
                svar = self.client.post(reverse("gennemse", args=[k.pk]), data, follow=True)
                self.assertEqual(svar.status_code, 200)
                self.assertRedirects(svar, reverse("gennemse", args=[k.pk]))
                self.assertContains(svar, "Vælg en medarbejder i listen.")
                self.assertContains(svar, "Bekræft match")
                l.refresh_from_db()
                self.assertIsNone(l.borger)
                self.assertNotEqual(l.match_status, Linje.Match.MATCHET)
                self.assertEqual(AuditLog.objects.count(), antal_logs)

    def test_beregn_afvises_naar_zone_mangler(self):
        k = self.opret_og_upload([("Person 3", "44:12")])
        svar = self.undtagelser(k)
        self.assertEqual(svar.status_code, 200)
        self.assertContains(svar, "Mangler zone")
        self.assertEqual(k.beregninger.count(), 0)

    def test_afvist_beregning_markerer_foerste_manglende_raekke(self):
        Borger.objects.create(navn="Person 4", tj_nr="9004", cpr="000000-0004", sats=Decimal("20"))
        k = self.opret_og_upload([("Person 1", "44:12"), ("Person 3", "10:00"), ("Person 4", "5:00")])
        svar = self.undtagelser(k)
        self.assertEqual(svar.status_code, 200)
        self.assertContains(svar, 'data-afvist="1"')
        self.assertContains(svar, "row g3 mangler", count=1)
        html = svar.content.decode()
        markeret = html.index("row g3 mangler")
        self.assertLess(html.index("Person 1"), markeret)
        self.assertLess(markeret, html.index("Person 3"))
        self.assertLess(markeret, html.index("Person 4"))
        self.assertContains(svar, "Beregningen blev ikke kørt, fordi nogle medarbejdere mangler oplysninger.")
        self.assertContains(svar, "1 af 3 klar")

    def test_undtagelser_uden_afvisning_har_ingen_markering(self):
        k = self.opret_og_upload([("Person 3", "10:00")])
        svar = self.client.get(reverse("undtagelser", args=[k.pk]))
        self.assertNotContains(svar, "data-afvist")
        self.assertNotContains(svar, "row g3 mangler")

    def test_gem_kladde_med_manglende_raekke_markerer_ikke(self):
        k = self.opret_og_upload([("Person 3", "10:00")])
        svar = self.undtagelser(k, handling="gem")
        self.assertRedirects(svar, reverse("undtagelser", args=[k.pk]))
        svar = self.client.get(reverse("undtagelser", args=[k.pk]))
        self.assertNotContains(svar, "data-afvist")
        self.assertNotContains(svar, "row g3 mangler")

    def test_hele_flowet_beregn_pdf_godkend_eksport(self):
        k = self.opret_og_upload([("Person 1", "44:12"), ("Person 2", "10:00")])
        svar = self.undtagelser(k)
        self.assertRedirects(svar, reverse("godkend", args=[k.pk]))
        b = k.beregninger.get()

        pdf = b.pdf.read()
        self.assertTrue(pdf.startswith(b"%PDF"))
        self.assertEqual(hashlib.sha256(pdf).hexdigest(), b.pdf_sha256)
        self.assertEqual(hashlib.sha256(lav_rapport(b.data)).hexdigest(), b.pdf_sha256)
        self.assertNotIn("000000-0001", b.snapshot)

        r1 = b.resultater.get(navn="Person 1")
        self.assertEqual(r1.loen_foer, Decimal("884.00"))
        self.assertEqual(r1.loen_efter, Decimal("531.00"))
        self.assertEqual(r1.udbetaling_991, Decimal("58.20"))

        spaerret = self.client.get(reverse("eksport", args=[k.pk]))
        self.assertRedirects(spaerret, reverse("godkend", args=[k.pk]))

        self.client.post(reverse("godkend", args=[k.pk]), {"handling": "godkend", "resultat": r1.pk})
        self.assertEqual(self.client.get(reverse("eksport", args=[k.pk])).status_code, 302)

        self.client.post(reverse("godkend", args=[k.pk]), {"handling": "godkend_alle"})
        k.refresh_from_db()
        self.assertEqual(k.status, Koersel.Status.GODKENDT)
        svar = self.client.get(reverse("eksport", args=[k.pk]))
        self.assertEqual(svar.status_code, 200)
        self.assertIn("Person 1", svar.content.decode("utf-8-sig"))

        self.assertEqual(self.client.get(reverse("rapport", args=[b.pk])).status_code, 200)

    def test_godkend_skaermen_har_ikke_rapportlinje_forklaring_eller_totaler(self):
        k = self.opret_og_upload([("Person 1", "44:12"), ("Person 2", "10:00")])
        self.undtagelser(k)
        b = k.beregninger.get()
        url = reverse("godkend", args=[k.pk])

        def tjek_fjernet(svar):
            self.assertEqual(svar.status_code, 200)
            # Linjen med rapportnummer, tidspunkt og "x af y godkendt"
            self.assertNotRegex(svar.content.decode(), r"\d+ af \d+ godkendt")
            self.assertNotContains(svar, f"Rapport {b.rapport_id} ·")
            # Sætningen under overskriften
            self.assertNotContains(svar, "Gennemgå, hvordan beløbet er beregnet")
            # Rækken med totaler
            self.assertNotContains(svar, 'class="tiles"')
            self.assertNotContains(svar, 'class="tile"')
            # Etiketterne må ikke stå som tile-etiketter; teksterne står i øvrigt
            # i beregningsgrundlaget pr. medarbejder, som nu har sin egen side.
            for etiket in ("Medarbejdere", "Løn før skat", "Løn efter skat", "Træk (lønart 991)", "Udbetaling (lønart 991)"):
                self.assertNotContains(svar, f'<div class="l">{etiket}</div>')
            # Det der bliver
            self.assertContains(svar, f"Godkend beregning · {k.navn}")
            self.assertContains(svar, f"Åbn PDF-rapport {b.rapport_id}")
            self.assertContains(svar, "Beregning pr. medarbejder")
            self.assertContains(svar, "Tidligere beregninger for måneden")

        svar = self.client.get(url)
        tjek_fjernet(svar)
        self.assertContains(svar, "Godkend alle")
        self.assertNotContains(svar, "Eksport er låst")

        self.client.post(url, {"handling": "godkend_alle"})
        svar = self.client.get(url)
        tjek_fjernet(svar)
        self.assertContains(svar, "Genåbn til rettelser")
        self.assertContains(svar, "Eksportér lønudtræk (CSV)")
        self.assertNotContains(svar, "Eksport er låst")

    def test_knapperne_oeverst_staar_i_ny_raekkefoelge_og_er_laast_til_alle_er_godkendt(self):
        k = self.opret_og_upload([("Person 1", "44:12"), ("Person 2", "10:00")])
        self.undtagelser(k)
        b = k.beregninger.get()
        r1 = b.resultater.order_by("pk").first()
        url = reverse("godkend", args=[k.pk])
        forklaring = "Kan først bruges, når alle medarbejdere er godkendt."
        pdf_tekst = f"Åbn PDF-rapport {b.rapport_id}"
        eksport_url = reverse("eksport", args=[k.pk])
        rapport_url = reverse("rapport", args=[b.pk])

        def top(html):
            start = html.index('<div class="actions"')
            return html[start:html.index("<h2>Beregning pr. medarbejder</h2>", start)]

        def laast(top_html, tekst):
            # En deaktiveret knap uden adresse, med forklaring til både mus og skærmlæser
            m = re.search(r"<button[^>]*>" + re.escape(tekst) + "</button>", top_html)
            self.assertIsNotNone(m, f"Der er ingen knap med teksten {tekst}")
            knap = m.group(0)
            self.assertIn("disabled", knap)
            self.assertNotIn("href", knap)
            self.assertNotIn('type="submit"', knap)
            self.assertIn(f'title="{forklaring}"', knap)
            self.assertIn(f'aria-label="{forklaring}"', knap)
            self.assertNotIn(f'>{tekst}</a>', top_html)

        def link(top_html, tekst, adresse):
            self.assertRegex(top_html, r'<a class="btn" href="' + re.escape(adresse) + r'"[^>]*>' + re.escape(tekst) + "</a>")
            self.assertNotRegex(top_html, r"<button[^>]*>" + re.escape(tekst) + "</button>")

        def raekkefoelge(top_html):
            positioner = [top_html.index(t) for t in ("Godkend alle", "Genåbn til rettelser", "Eksportér lønudtræk (CSV)", pdf_tekst)]
            self.assertEqual(positioner, sorted(positioner))

        def genaabn_er_aktiv(top_html):
            m = re.search(r"<form[^>]*>(?:(?!</form>).)*value=\"genaabn\"(?:(?!</form>).)*</form>", top_html, re.S)
            self.assertIsNotNone(m)
            self.assertIn(">Genåbn til rettelser</button>", m.group(0))
            self.assertNotIn("disabled", m.group(0))

        # Ingen er godkendt
        html = self.client.get(url).content.decode()
        t = top(html)
        raekkefoelge(t)
        self.assertIn('value="godkend_alle"', t)
        self.assertNotIn('class="godkendt"', t)
        self.assertIn(">Godkend alle</button>", t)
        genaabn_er_aktiv(t)
        laast(t, "Eksportér lønudtræk (CSV)")
        laast(t, pdf_tekst)
        self.assertNotIn(eksport_url, t)
        self.assertNotIn(rapport_url, t)
        self.assertNotIn("Eksport er låst", html)
        self.assertNotIn('<span class="badge ok">Godkendt</span>', html)

        # Eksportadressen afviser stadig på serveren
        self.assertRedirects(self.client.get(eksport_url), url)

        # Kun den ene godkendt: stadig låst
        self.client.post(url, {"handling": "godkend", "resultat": r1.pk})
        t = top(self.client.get(url).content.decode())
        laast(t, "Eksportér lønudtræk (CSV)")
        laast(t, pdf_tekst)
        self.assertIn('value="godkend_alle"', t)
        self.assertRedirects(self.client.get(eksport_url), url)

        # Alle godkendt
        self.client.post(url, {"handling": "godkend_alle"})
        html = self.client.get(url).content.decode()
        t = top(html)
        raekkefoelge(t)
        link(t, "Eksportér lønudtræk (CSV)", eksport_url)
        self.assertRegex(t, r'<a class="btn" href="' + re.escape(rapport_url) + r'" target="_blank" rel="noopener">' + re.escape(pdf_tekst) + "</a>")
        self.assertNotRegex(t, r"<button[^>]*>" + re.escape(pdf_tekst) + "</button>")
        self.assertNotIn('value="godkend_alle"', t)
        m = re.search(r"<button[^>]*>Godkend alle</button>", t)
        self.assertIsNotNone(m)
        self.assertIn('type="button" disabled class="godkendt"', m.group(0))
        self.assertIn('title="Alle medarbejdere er godkendt."', m.group(0))
        self.assertIn('aria-label="Alle medarbejdere er godkendt."', m.group(0))
        genaabn_er_aktiv(t)
        self.assertNotIn("Eksport er låst", html)
        self.assertNotIn('<span class="badge ok">Godkendt</span>', html)
        self.assertNotIn(forklaring, t)
        self.assertEqual(self.client.get(eksport_url).status_code, 200)

        # Genåbn: tilbage til den låste tilstand, og knappen er der hele vejen
        self.client.post(url, {"handling": "genaabn"})
        t = top(self.client.get(url).content.decode())
        laast(t, "Eksportér lønudtræk (CSV)")
        laast(t, pdf_tekst)
        genaabn_er_aktiv(t)
        self.assertRedirects(self.client.get(eksport_url), url)

    def _raekker_og_top(self, url):
        html = self.client.get(url).content.decode()
        start = html.index('<div class="actions"')
        midt = html.index("<h2>Beregning pr. medarbejder</h2>", start)
        slut = html.index("<h2>Tidligere beregninger for måneden</h2>", midt)
        return html[start:midt], html[midt:slut]

    def test_genaabn_fra_helt_godkendt_koersel_fjerner_alle_godkendelser_og_laaser_eksport_og_pdf(self):
        k = self.opret_og_upload([("Person 1", "44:12"), ("Person 2", "10:00")])
        self.undtagelser(k)
        url = reverse("godkend", args=[k.pk])
        b = k.aktiv_beregning
        self.client.post(url, {"handling": "godkend_alle"})
        self.assertEqual(b.resultater.filter(godkendt=True).count(), 2)
        k.refresh_from_db()
        self.assertEqual(k.status, Koersel.Status.GODKENDT)

        svar = self.client.post(url, {"handling": "genaabn"}, follow=True)
        self.assertRedirects(svar, url)
        self.assertContains(svar, "Kørslen er genåbnet. 2 godkendelser er fjernet. Godkend igen, eller ret oplysningerne under Undtagelser og kør en ny beregning.")
        # Siden man lander på, er Godkend-skærmen selv: ingen grønne rækkeknapper, Godkend alle aktiv, PDF og eksport låst
        html = svar.content.decode()
        start = html.index('<div class="actions"')
        midt = html.index("<h2>Beregning pr. medarbejder</h2>", start)
        slut = html.index("<h2>Tidligere beregninger for måneden</h2>", midt)
        self.assertNotIn('class="godkendt"', html[midt:slut])
        self.assertNotIn('class="godkendt"', html[start:midt])
        self.assertIn('value="godkend_alle"', html[start:midt])
        self.assertIn("Kan først bruges, når alle medarbejdere er godkendt.", html[start:midt])
        self.assertNotIn(reverse("eksport", args=[k.pk]), html[start:midt])

        k.refresh_from_db()
        self.assertEqual(k.status, Koersel.Status.BEREGNET)
        self.assertEqual(b.resultater.count(), 2)
        for r in b.resultater.all():
            self.assertFalse(r.godkendt)
            self.assertIsNone(r.godkendt_af)
            self.assertIsNone(r.godkendt_tidspunkt)
        self.assertFalse(services.kan_eksportere(k))
        # Samme beregning, ingen ny
        self.assertEqual(k.beregninger.count(), 1)
        self.assertRedirects(self.client.get(reverse("eksport", args=[k.pk])), url)

        # Godkend-skærmen: ingen grønne rækkeknapper, Godkend alle er aktiv, PDF og eksport er låst
        top, raekker = self._raekker_og_top(url)
        self.assertNotIn('class="godkendt"', raekker)
        self.assertEqual(raekker.count('<button type="submit">Godkend</button>'), 2)
        self.assertIn('value="godkend_alle"', top)
        self.assertNotIn('class="godkendt"', top)
        self.assertNotRegex(top, r'<a class="btn" href="' + re.escape(reverse("eksport", args=[k.pk])))
        self.assertNotRegex(top, r'<a class="btn" href="' + re.escape(reverse("rapport", args=[b.pk])))
        self.assertIn("Kan først bruges, når alle medarbejdere er godkendt.", top)

        # Man kan godkende igen uden ny beregning
        self.client.post(url, {"handling": "godkend_alle"})
        k.refresh_from_db()
        self.assertTrue(services.kan_eksportere(k))

    def test_genaabn_fra_delvist_godkendt_koersel_fjerner_de_givne_godkendelser(self):
        k = self.opret_og_upload([("Person 1", "44:12"), ("Person 2", "10:00")])
        self.undtagelser(k)
        url = reverse("godkend", args=[k.pk])
        b = k.aktiv_beregning
        self.client.post(url, {"handling": "godkend", "resultat": b.resultater.order_by("pk").first().pk})
        self.assertEqual(b.resultater.filter(godkendt=True).count(), 1)
        k.refresh_from_db()
        self.assertEqual(k.status, Koersel.Status.BEREGNET)

        svar = self.client.post(url, {"handling": "genaabn"}, follow=True)
        self.assertRedirects(svar, url)
        self.assertContains(svar, "Kørslen er genåbnet. 1 godkendelse er fjernet. Godkend igen, eller ret oplysningerne under Undtagelser og kør en ny beregning.")
        self.assertEqual(b.resultater.filter(godkendt=True).count(), 0)
        self.assertEqual(b.resultater.filter(godkendt_af__isnull=False).count(), 0)
        self.assertEqual(b.resultater.filter(godkendt_tidspunkt__isnull=False).count(), 0)
        k.refresh_from_db()
        self.assertEqual(k.status, Koersel.Status.BEREGNET)

        top, raekker = self._raekker_og_top(url)
        self.assertNotIn('class="godkendt"', raekker)
        self.assertEqual(raekker.count('<button type="submit">Godkend</button>'), 2)
        self.assertIn('value="godkend_alle"', top)

    def test_genaabn_skriver_antal_fjernede_godkendelser_i_auditloggen(self):
        k = self.opret_og_upload([("Person 1", "44:12"), ("Person 2", "10:00")])
        self.undtagelser(k)
        url = reverse("godkend", args=[k.pk])
        self.client.post(url, {"handling": "godkend_alle"})
        self.client.post(url, {"handling": "genaabn"})
        post = AuditLog.objects.get(handling="Genåbnede kørsel")
        self.assertEqual(post.bruger, self.bruger)
        self.assertEqual(post.objekt_type, "Koersel")
        self.assertEqual(post.objekt_id, str(k.pk))
        self.assertIn("fjernede godkendelser: 2", post.detaljer)

        # Delvist godkendt: også med antallet
        self.undtagelser(k)
        b = k.aktiv_beregning
        self.client.post(url, {"handling": "godkend", "resultat": b.resultater.order_by("pk").first().pk})
        self.client.post(url, {"handling": "genaabn"})
        seneste = AuditLog.objects.filter(handling="Genåbnede kørsel").first()
        self.assertIn("fjernede godkendelser: 1", seneste.detaljer)

    def test_genaabn_uden_godkendelser_aendrer_ingenting_og_giver_ingen_fejl(self):
        k = self.opret_og_upload([("Person 1", "44:12"), ("Person 2", "10:00")])
        self.undtagelser(k)
        url = reverse("godkend", args=[k.pk])
        b = k.aktiv_beregning
        antal_logs = AuditLog.objects.count()
        svar = self.client.post(url, {"handling": "genaabn"}, follow=True)
        self.assertRedirects(svar, url)
        self.assertNotContains(svar, "genåbnet")
        k.refresh_from_db()
        self.assertEqual(k.status, Koersel.Status.BEREGNET)
        self.assertEqual(b.resultater.filter(godkendt=False).count(), 2)
        self.assertEqual(AuditLog.objects.count(), antal_logs)
        self.assertFalse(AuditLog.objects.filter(handling="Genåbnede kørsel").exists())

    def test_genaabn_foer_en_beregning_er_koert_giver_ingen_fejl(self):
        k = self.opret_og_upload([("Person 1", "44:12")])
        svar = self.client.post(reverse("godkend", args=[k.pk]), {"handling": "genaabn"})
        self.assertRedirects(svar, reverse("undtagelser", args=[k.pk]))
        self.assertFalse(AuditLog.objects.filter(handling="Genåbnede kørsel").exists())

    def test_genaabn_kraever_login(self):
        k = self.opret_og_upload([("Person 1", "44:12")])
        self.undtagelser(k)
        url = reverse("godkend", args=[k.pk])
        self.client.post(url, {"handling": "godkend_alle"})
        self.client.logout()
        svar = self.client.post(url, {"handling": "genaabn"})
        self.assertEqual(svar.status_code, 302)
        self.assertIn("/konto/login/", svar["Location"])
        # Intet er fjernet
        self.assertEqual(k.aktiv_beregning.resultater.filter(godkendt=True).count(), 1)
        k.refresh_from_db()
        self.assertEqual(k.status, Koersel.Status.GODKENDT)

    def test_godkend_skaermen_har_ikke_knappen_tilbage_til_undtagelser_men_trinsoejlen_linker_dertil(self):
        k = self.opret_og_upload([("Person 1", "44:12"), ("Person 2", "10:00")])
        self.undtagelser(k)
        url = reverse("godkend", args=[k.pk])
        undtagelser_url = reverse("undtagelser", args=[k.pk])

        def tjek(svar):
            self.assertEqual(svar.status_code, 200)
            self.assertNotContains(svar, "Tilbage til undtagelser")
            # Trinsøjlen har stadig sit eget link til Undtagelser, og det er det eneste
            self.assertContains(
                svar,
                f'<a class="stp done" href="{undtagelser_url}">3 Undtagelser</a>',
                html=False,
            )
            self.assertEqual(svar.content.decode().count(f'href="{undtagelser_url}"'), 1)

        tjek(self.client.get(url))
        self.client.post(url, {"handling": "godkend_alle"})
        tjek(self.client.get(url))

    def test_godkend_skaermen_har_knap_til_beregningsgrundlag_i_stedet_for_foldet_tabel(self):
        k = self.opret_og_upload([("Person 1", "44:12"), ("Person 2", "10:00")])
        self.undtagelser(k)
        b = k.beregninger.get()
        url = reverse("godkend", args=[k.pk])

        def tjek(svar):
            self.assertEqual(svar.status_code, 200)
            self.assertNotContains(svar, "Se beregningsgrundlag")
            self.assertNotContains(svar, "<details")
            for r in b.resultater.all():
                self.assertContains(
                    svar,
                    f'<a class="btn" href="{reverse("grundlag", args=[k.pk, r.pk])}" target="_blank" rel="noopener">Beregningsgrundlag</a>',
                    html=False,
                )
            self.assertEqual(svar.content.decode().count(">Beregningsgrundlag</a>"), 2)

        tjek(self.client.get(url))
        self.client.post(url, {"handling": "godkend_alle"})
        svar = self.client.get(url)
        tjek(svar)
        # Knappen står stadig, også når borgeren er godkendt
        self.assertContains(svar, "Godkendt")

    def test_godkend_og_beregningsgrundlag_staar_i_samme_celle_med_godkend_foerst(self):
        k = self.opret_og_upload([("Person 1", "44:12"), ("Person 2", "10:00")])
        self.undtagelser(k)
        url = reverse("godkend", args=[k.pk])

        def celler(svar):
            html = svar.content.decode()
            # Hver celle går fra sin åbning til næste celle (eller rækkens slutning)
            dele = html.split('<div class="knapper">')[1:]
            return [d.split("</div>")[0] for d in dele]

        # Ikke godkendt: begge knapper i samme celle, og "Beregningsgrundlag" efter "Godkend"
        udfald = celler(self.client.get(url))
        self.assertEqual(len(udfald), 2)
        for celle in udfald:
            self.assertEqual(celle.count(">Godkend</button>"), 1)
            self.assertEqual(celle.count(">Beregningsgrundlag</a>"), 1)
            self.assertLess(celle.index(">Godkend</button>"), celle.index(">Beregningsgrundlag</a>"))

        # Godkendt: cellen står stadig, men "Godkend" er en deaktiveret knap uden formular
        self.client.post(url, {"handling": "godkend_alle"})
        udfald = celler(self.client.get(url))
        self.assertEqual(len(udfald), 2)
        for celle in udfald:
            self.assertEqual(celle.count(">Godkend</button>"), 1)
            self.assertIn("disabled", celle)
            self.assertNotIn("<form", celle)
            self.assertEqual(celle.count(">Beregningsgrundlag</a>"), 1)
            self.assertLess(celle.index(">Godkend</button>"), celle.index(">Beregningsgrundlag</a>"))

    def test_godkendte_og_ikke_godkendte_raekker_har_samme_opbygning_uden_maerker(self):
        k = self.opret_og_upload([("Person 1", "44:12"), ("Person 2", "10:00")])
        self.undtagelser(k)
        b = k.beregninger.get()
        r1 = b.resultater.order_by("pk").first()
        url = reverse("godkend", args=[k.pk])

        def celler(html):
            return [d.split("</div>")[0] for d in html.split('<div class="knapper">')[1:]]

        def raekker(html):
            # Alt i en række før knapcellen: navn, tre beløb og eventuelle mærker
            return [d.split('<div class="knapper">')[0] for d in html.split('<div class="row g4"')[1:]]

        def toplevel(celle):
            # De to elementer i cellen: formularen eller den deaktiverede knap, så linket
            return re.findall(r'<(form|button type="button" disabled class="godkendt"|a class="btn")', celle)

        def uden_maerker(html):
            # Hverken "Afventer" eller "Godkendt" som mærke, og ingen mærker overhovedet i rækkerne
            self.assertNotIn("Afventer", html)
            self.assertEqual(len(raekker(html)), 2)
            for raekke in raekker(html):
                self.assertNotIn("badge", raekke)
                self.assertNotIn("Afventer", raekke)
                self.assertNotIn("Godkendt", raekke)

        # Ingen række har mærker, og ingen grøn knap står før nogen er godkendt
        svar = self.client.get(url)
        uden_maerker(svar.content.decode())
        self.assertNotContains(svar, 'class="godkendt"')

        # Kun den ene medarbejder godkendes
        self.client.post(url, {"handling": "godkend", "resultat": r1.pk})
        svar = self.client.get(url)
        html = svar.content.decode()
        uden_maerker(html)
        # Der er intet samlet mærke øverst
        self.assertNotIn('<span class="badge ok">Godkendt</span>', html)

        # Rækkefølgen af rækkerne er ikke afgørende; den godkendte er den uden formular
        a, b_celle = celler(html)
        godkendt, ikke_godkendt = (b_celle, a) if "<form" in a else (a, b_celle)
        # Begge celler har to elementer i samme rækkefølge: Godkend, så "Beregningsgrundlag"
        self.assertEqual(toplevel(ikke_godkendt), ["form", 'a class="btn"'])
        self.assertEqual(toplevel(godkendt), ['button type="button" disabled class="godkendt"', 'a class="btn"'])
        # Den ikke-godkendte række har den aktive formular og ingen grøn knap
        self.assertIn(">Godkend</button>", ikke_godkendt)
        self.assertIn('name="handling" value="godkend"', ikke_godkendt)
        self.assertNotIn("godkendt", ikke_godkendt)
        self.assertNotIn("disabled", ikke_godkendt)
        # Den godkendte række: knappen hedder stadig "Godkend", kan ikke trykkes, og er ikke en formular
        self.assertEqual(godkendt.count(">Godkend</button>"), 1)
        self.assertNotIn("<form", godkendt)
        self.assertNotIn('type="submit"', godkendt)
        # Hvem og hvornår står som oplysning på knappen, til både mus og skærmlæser
        tidspunkt = timezone.localtime(b.resultater.get(pk=r1.pk).godkendt_tidspunkt).strftime("%d.%m.%Y %H:%M")
        oplysning = f"Godkendt af tester {tidspunkt}"
        self.assertIn(f'title="{oplysning}"', godkendt)
        self.assertIn(f'aria-label="{oplysning}"', godkendt)

        # Efter "Godkend alle" har begge rækker samme opbygning og ingen mærker
        self.client.post(url, {"handling": "godkend_alle"})
        svar = self.client.get(url)
        html = svar.content.decode()
        uden_maerker(html)
        # Mærket "Godkendt" øverst er væk; den grønne "Godkend alle" viser oplysningen
        self.assertNotIn('<span class="badge ok">Godkendt</span>', html)
        self.assertNotIn('class="badge ok"', html)
        for celle in celler(html):
            self.assertEqual(toplevel(celle), ['button type="button" disabled class="godkendt"', 'a class="btn"'])
            self.assertIn('title="Godkendt af tester ', celle)
            self.assertIn('aria-label="Godkendt af tester ', celle)

    def test_den_groenne_godkend_knap_har_egen_farve_uden_hover_og_pladsholderen_er_vaek(self):
        css = (settings.BASE_DIR / "static" / "css" / "app.css").read_text(encoding="utf-8")
        self.assertNotIn(".plads", css)
        self.assertNotIn("plads", (settings.BASE_DIR / "templates" / "koersel" / "godkend.html").read_text(encoding="utf-8"))
        regel = re.search(r"^(\.knapper button\.godkendt[^{]*)\{([^}]*)\}", css, re.M)
        self.assertIsNotNone(regel)
        vaelgere, indhold = regel.groups()
        self.assertIn("background: var(--ok-bg)", indhold)
        self.assertIn("color: var(--ok)", indhold)
        self.assertIn("cursor: default", indhold)
        # Hover og deaktiveret tilstand får samme farver, så hverken hover eller browserens grå tekst slår igennem
        self.assertIn(".knapper button.godkendt:hover", vaelgere)
        self.assertIn(".knapper button.godkendt:disabled", vaelgere)
        # Mørk tilstand bruger de samme variabler, som er defineret dér
        self.assertRegex(css, r"prefers-color-scheme: dark[^}]*--ok-bg:[^;]*; --ok:")
        # Samme padding og kantbredde som den almindelige knap: reglen sætter ingen af dem
        self.assertNotIn("padding", indhold)
        self.assertNotIn("border-width", indhold)

    def test_knapkolonnen_har_ingen_fast_bredde_og_knapperne_ombrydes_ikke(self):
        css = (settings.BASE_DIR / "static" / "css" / "app.css").read_text(encoding="utf-8")

        def regel(vaelger):
            m = re.search(r"^" + re.escape(vaelger) + r" \{([^}]*)\}", css, re.M)
            self.assertIsNotNone(m, f"Reglen {vaelger} findes ikke")
            return m.group(1)

        g4 = regel(".g4")
        kolonner = re.findall(r"minmax\([^)]*\)|\S+", re.search(r"grid-template-columns:\s*([^;]+);", g4).group(1))
        # Fem kolonner: navn, tre beløb og knapper; statuskolonnen findes ikke længere
        self.assertEqual(len(kolonner), 5)
        # Sidste kolonne er knapkolonnen: den fylder det knapperne selv fylder, ikke en fast pixelbredde
        self.assertEqual(kolonner[-1], "max-content")
        # Mere luft mellem kolonnerne end de tidligere 10 px
        self.assertGreaterEqual(int(re.search(r"column-gap:\s*(\d+)px", g4).group(1)), 20)
        knapper = regel(".knapper")
        self.assertIn("flex-wrap: nowrap", knapper)
        self.assertIn("white-space: nowrap", knapper)
        # Cellen har ingen fast bredde i skabelonen
        self.assertNotIn('class="knapper" style', (settings.BASE_DIR / "templates" / "koersel" / "godkend.html").read_text(encoding="utf-8"))

    def test_beregningsgrundlag_viser_borger_og_tabel_uden_cpr_og_adresse(self):
        self.p1.adresse = "Fiktiv Vej 1, 9000 Testby"
        self.p1.save()
        k = self.opret_og_upload([("Person 1", "44:12"), ("Person 2", "10:00")])
        self.undtagelser(k)
        b = k.beregninger.get()
        r = b.resultater.get(navn="Person 1")
        svar = self.client.get(reverse("grundlag", args=[k.pk, r.pk]))
        self.assertEqual(svar.status_code, 200)
        self.assertContains(svar, "Beregningsgrundlag · Person 1")
        self.assertContains(svar, "9001")
        self.assertContains(svar, "44:12")
        self.assertContains(svar, "Kører selv")
        for kolonne in ("Nr.", "Trin", "Formel", "Indsat", "Resultat"):
            self.assertContains(svar, f"<th>{kolonne}</th>" if kolonne != "Resultat" else '<th class="r">Resultat</th>')
        for trin in r.detaljer["grundlag"]:
            self.assertContains(svar, trin["titel"])
            if trin.get("note"):
                self.assertContains(svar, trin["note"])
        self.assertNotContains(svar, "000000-0001")
        self.assertNotContains(svar, "Fiktiv Vej")
        # Ingen trinsøjle og ingen godkendelsesknap
        self.assertNotContains(svar, "med-trin")
        self.assertNotContains(svar, 'name="handling"')

    def test_beregningsgrundlag_giver_ikke_fundet_for_ugyldigt_nummer_eller_fremmed_borger(self):
        k = self.opret_og_upload([("Person 1", "44:12")])
        # Ingen beregning endnu
        self.assertEqual(self.client.get(reverse("grundlag", args=[k.pk, 1])).status_code, 404)
        self.undtagelser(k)
        r = k.aktiv_beregning.resultater.get()
        self.assertEqual(self.client.get(reverse("grundlag", args=[k.pk, r.pk + 999])).status_code, 404)
        self.assertEqual(self.client.get(reverse("grundlag", args=[k.pk + 999, r.pk])).status_code, 404)
        # Et resultat fra en tidligere beregning hører ikke til den gældende
        self.client.post(reverse("godkend", args=[k.pk]), {"handling": "godkend_alle"})
        self.client.post(reverse("godkend", args=[k.pk]), {"handling": "genaabn"})
        self.undtagelser(k)
        self.assertEqual(self.client.get(reverse("grundlag", args=[k.pk, r.pk])).status_code, 404)
        ny = k.aktiv_beregning.resultater.get()
        self.assertEqual(self.client.get(reverse("grundlag", args=[k.pk, ny.pk])).status_code, 200)

    def test_beregningsgrundlag_kraever_login(self):
        k = self.opret_og_upload([("Person 1", "44:12")])
        self.undtagelser(k)
        r = k.aktiv_beregning.resultater.get()
        self.client.logout()
        svar = self.client.get(reverse("grundlag", args=[k.pk, r.pk]))
        self.assertEqual(svar.status_code, 302)
        self.assertIn("/konto/login/", svar["Location"])

    def test_ny_beregning_ophaever_tidligere_godkendelser(self):
        k = self.opret_og_upload([("Person 1", "44:12")])
        self.undtagelser(k)
        self.client.post(reverse("godkend", args=[k.pk]), {"handling": "godkend_alle"})
        self.client.post(reverse("godkend", args=[k.pk]), {"handling": "genaabn"})
        self.undtagelser(k)
        self.assertEqual(k.beregninger.count(), 2)
        ny = k.aktiv_beregning
        self.assertEqual(ny.nummer, 2)
        self.assertFalse(ny.alle_godkendt)
        k.refresh_from_db()
        self.assertEqual(k.status, Koersel.Status.BEREGNET)
        self.assertEqual(self.client.get(reverse("eksport", args=[k.pk])).status_code, 302)
        self.assertTrue(k.beregninger.get(nummer=1).pdf)

    FORAELDET = "Beregningen er forældet."
    FORAELDET_BESKED = "Beregningen er forældet. Der er ændret noget, siden den blev kørt. Kør Beregn igen under Undtagelser, før du kan godkende."
    FORAELDET_KNAP = "Beregningen er forældet. Kør Beregn igen."

    def _beregnet(self):
        k = self.opret_og_upload([("Person 1", "44:12"), ("Person 2", "10:00")])
        self.undtagelser(k)
        k.refresh_from_db()
        self.assertEqual(k.status, Koersel.Status.BEREGNET)
        return k

    def _gem_og_fortsaet(self, k):
        data = {"handling": "gem"}
        for l in k.linjer.select_related("borger"):
            data.update({f"afstand_{l.pk}": str(l.borger.afstand_km), f"transport_{l.pk}": l.transport})
        return self.client.post(reverse("gennemse", args=[k.pk]), data)

    def _top_og_raekker(self, k):
        return self._raekker_og_top(reverse("godkend", args=[k.pk]))

    def _tjek_foraeldet_skaerm(self, k):
        url = reverse("godkend", args=[k.pk])
        b = k.aktiv_beregning
        svar = self.client.get(url)
        self.assertContains(svar, self.FORAELDET)
        self.assertContains(svar, "Der er ændret noget, siden den blev kørt. Kør Beregn igen under")
        self.assertContains(svar, f'<a href="{reverse("undtagelser", args=[k.pk])}">Undtagelser</a>')
        top, raekker = self._top_og_raekker(k)
        # Godkend alle og alle rækkeknapper: ikke trykbare, ikke grønne, med forklaring
        self.assertNotIn('value="godkend_alle"', top)
        self.assertNotIn('class="godkendt"', top)
        self.assertNotIn('class="godkendt"', raekker)
        self.assertNotIn('value="godkend"', raekker)
        self.assertEqual(raekker.count(f'<button type="button" disabled title="{self.FORAELDET_KNAP}" aria-label="{self.FORAELDET_KNAP}">Godkend</button>'), 2)
        self.assertIn(f'<button type="button" disabled title="{self.FORAELDET_KNAP}" aria-label="{self.FORAELDET_KNAP}">Godkend alle</button>', top)
        # Eksport og PDF: ikke trykbare, ingen adresser
        for tekst in ("Eksportér lønudtræk (CSV)", f"Åbn PDF-rapport {b.rapport_id}"):
            self.assertIn(f'<button type="button" disabled title="{self.FORAELDET_KNAP}" aria-label="{self.FORAELDET_KNAP}">{tekst}</button>', top)
        self.assertNotIn(reverse("eksport", args=[k.pk]), top)
        self.assertNotIn(reverse("rapport", args=[b.pk]), top)
        # Genåbn og Beregningsgrundlag kan stadig trykkes
        self.assertRegex(top, r'value="genaabn".*?<button type="submit">Genåbn til rettelser</button>')
        for r in b.resultater.all():
            self.assertIn(f'<a class="btn" href="{reverse("grundlag", args=[k.pk, r.pk])}"', raekker)

    def _tjek_godkendelse_afvist(self, k):
        url = reverse("godkend", args=[k.pk])
        b = k.aktiv_beregning
        foer_status = Koersel.objects.get(pk=k.pk).status
        antal_logs = AuditLog.objects.filter(handling="Godkendte beregning").count()
        r1 = b.resultater.order_by("pk").first()
        for data in ({"handling": "godkend", "resultat": r1.pk}, {"handling": "godkend_alle"}, {"handling": "godkend", "resultat": ""}):
            svar = self.client.post(url, data, follow=True)
            self.assertRedirects(svar, url)
            self.assertContains(svar, self.FORAELDET_BESKED)
            self.assertNotContains(svar, "medarbejder(e) godkendt")
            self.assertEqual(b.resultater.filter(godkendt=True).count(), 0)
            self.assertEqual(Koersel.objects.get(pk=k.pk).status, foer_status)
        self.assertEqual(AuditLog.objects.filter(handling="Godkendte beregning").count(), antal_logs)
        self.assertFalse(services.kan_eksportere(k))
        self.assertRedirects(self.client.get(reverse("eksport", args=[k.pk])), url)
        self.assertEqual(k.aktiv_beregning.pk, b.pk)

    def test_ny_upload_efter_genaabning_goer_beregningen_foraeldet(self):
        k = self._beregnet()
        url = reverse("godkend", args=[k.pk])
        self.client.post(url, {"handling": "godkend_alle"})
        self.client.post(url, {"handling": "genaabn"})
        k.refresh_from_db()
        self.assertEqual(k.status, Koersel.Status.BEREGNET)
        self.assertFalse(k.beregning_foraeldet)

        self.client.post(reverse("upload", args=[k.pk]), {"fil": xlsx([("Person 1", "1:00"), ("Person 2", "2:00")], navn="ny.xlsx")})
        k.refresh_from_db()
        self.assertEqual(k.status, Koersel.Status.KLADDE)
        self.assertTrue(k.beregning_foraeldet)
        self._tjek_foraeldet_skaerm(k)
        self._tjek_godkendelse_afvist(k)

    def test_gem_kladde_efter_beregning_goer_beregningen_foraeldet(self):
        k = self._beregnet()
        svar = self.undtagelser(k, handling="gem")
        self.assertRedirects(svar, reverse("undtagelser", args=[k.pk]))
        k.refresh_from_db()
        self.assertEqual(k.status, Koersel.Status.KLADDE)
        self.assertTrue(k.beregning_foraeldet)
        self._tjek_foraeldet_skaerm(k)
        self._tjek_godkendelse_afvist(k)

    def test_gem_og_fortsaet_efter_beregning_goer_beregningen_foraeldet(self):
        k = self._beregnet()
        svar = self._gem_og_fortsaet(k)
        self.assertRedirects(svar, reverse("undtagelser", args=[k.pk]))
        k.refresh_from_db()
        self.assertEqual(k.status, Koersel.Status.KLADDE)
        self._tjek_foraeldet_skaerm(k)
        self._tjek_godkendelse_afvist(k)

    def test_fortsaet_uden_at_gemme_goer_ikke_beregningen_foraeldet(self):
        k = self._beregnet()
        svar = self.client.post(reverse("gennemse", args=[k.pk]), {"handling": "fortsaet"})
        self.assertRedirects(svar, reverse("undtagelser", args=[k.pk]))
        k.refresh_from_db()
        self.assertEqual(k.status, Koersel.Status.BEREGNET)
        self.assertNotContains(self.client.get(reverse("godkend", args=[k.pk])), self.FORAELDET)

    def test_afvist_beregn_med_gemte_rettelser_goer_beregningen_foraeldet(self):
        k = self._beregnet()
        l = k.linjer.get(fil_navn="Person 1")
        svar = self.undtagelser(k, **{f"z_{l.pk}": ""})
        self.assertEqual(svar.status_code, 200)
        self.assertEqual(k.aktiv_beregning.nummer, 1)
        k.refresh_from_db()
        self.assertEqual(k.status, Koersel.Status.KLADDE)
        self._tjek_foraeldet_skaerm(k)
        self._tjek_godkendelse_afvist(k)

    def test_afklar_efter_beregning_goer_beregningen_foraeldet(self):
        k = self._beregnet()
        l = k.linjer.get(fil_navn="Person 1")
        self.client.post(reverse("gennemse", args=[k.pk]), {"handling": "afklar", "linje": l.pk, "borger": self.p3.pk})
        k.refresh_from_db()
        self.assertEqual(k.status, Koersel.Status.KLADDE)
        self._tjek_godkendelse_afvist(k)

    def test_afklar_med_ugyldigt_valg_aendrer_ikke_status(self):
        k = self._beregnet()
        l = k.linjer.get(fil_navn="Person 1")
        self.client.post(reverse("gennemse", args=[k.pk]), {"handling": "afklar", "linje": l.pk, "borger": ""})
        k.refresh_from_db()
        self.assertEqual(k.status, Koersel.Status.BEREGNET)

    def test_genaabn_og_beregningsgrundlag_virker_naar_beregningen_er_foraeldet(self):
        k = self._beregnet()
        url = reverse("godkend", args=[k.pk])
        self.client.post(url, {"handling": "godkend", "resultat": k.aktiv_beregning.resultater.order_by("pk").first().pk})
        self.undtagelser(k, handling="gem")
        r = k.aktiv_beregning.resultater.order_by("pk").first()
        self.assertEqual(self.client.get(reverse("grundlag", args=[k.pk, r.pk])).status_code, 200)
        svar = self.client.post(url, {"handling": "genaabn"}, follow=True)
        self.assertRedirects(svar, url)
        self.assertContains(svar, "Kørslen er genåbnet. 1 godkendelse er fjernet. Godkend igen, eller ret oplysningerne under Undtagelser og kør en ny beregning.")
        self.assertEqual(k.aktiv_beregning.resultater.filter(godkendt=True).count(), 0)
        k.refresh_from_db()
        self.assertEqual(k.status, Koersel.Status.KLADDE)
        self._tjek_foraeldet_skaerm(k)

    def test_ny_beregn_goer_en_foraeldet_beregning_aktuel_igen(self):
        k = self._beregnet()
        self.undtagelser(k, handling="gem")
        k.refresh_from_db()
        self.assertTrue(k.beregning_foraeldet)
        svar = self.undtagelser(k)
        self.assertRedirects(svar, reverse("godkend", args=[k.pk]))
        k.refresh_from_db()
        self.assertEqual(k.status, Koersel.Status.BEREGNET)
        self.assertFalse(k.beregning_foraeldet)
        self.assertEqual(k.aktiv_beregning.nummer, 2)
        url = reverse("godkend", args=[k.pk])
        svar = self.client.get(url)
        self.assertNotContains(svar, self.FORAELDET)
        top, raekker = self._top_og_raekker(k)
        self.assertIn('value="godkend_alle"', top)
        self.assertEqual(raekker.count('<button type="submit">Godkend</button>'), 2)
        self.assertIn("Kan først bruges, når alle medarbejdere er godkendt.", top)
        self.assertNotIn(self.FORAELDET_KNAP, top + raekker)
        # Godkend, og eksport og PDF åbner
        svar = self.client.post(url, {"handling": "godkend_alle"}, follow=True)
        self.assertContains(svar, "2 medarbejder(e) godkendt.")
        k.refresh_from_db()
        self.assertEqual(k.status, Koersel.Status.GODKENDT)
        self.assertTrue(services.kan_eksportere(k))
        self.assertEqual(self.client.get(reverse("eksport", args=[k.pk])).status_code, 200)
        self.assertEqual(self.client.get(reverse("rapport", args=[k.aktiv_beregning.pk])).status_code, 200)

    def test_en_koersel_gaar_aldrig_fra_kladde_til_godkendt(self):
        k = self._beregnet()
        self.undtagelser(k, handling="gem")
        b = k.aktiv_beregning
        k.refresh_from_db()
        self.assertEqual(k.status, Koersel.Status.KLADDE)
        with self.assertRaises(services.BeregningForaeldet):
            services.godkend(b, self.bruger)
        with self.assertRaises(services.BeregningForaeldet):
            services.godkend(b, self.bruger, b.resultater.first().pk)
        # Også når alle resultater allerede står som godkendt, skifter status ikke af sig selv
        b.resultater.update(godkendt=True)
        self.assertFalse(services.kan_eksportere(k))
        self.assertRedirects(self.client.get(reverse("eksport", args=[k.pk])), reverse("godkend", args=[k.pk]))
        self.client.post(reverse("godkend", args=[k.pk]), {"handling": "godkend_alle"})
        k.refresh_from_db()
        self.assertEqual(k.status, Koersel.Status.KLADDE)

    def test_ikke_foraeldet_beregning_virker_som_foer(self):
        k = self._beregnet()
        url = reverse("godkend", args=[k.pk])
        svar = self.client.get(url)
        self.assertNotContains(svar, self.FORAELDET)
        self.assertNotContains(svar, "forældet")
        svar = self.client.post(url, {"handling": "godkend", "resultat": k.aktiv_beregning.resultater.order_by("pk").first().pk}, follow=True)
        self.assertContains(svar, "1 medarbejder(e) godkendt.")
        svar = self.client.post(url, {"handling": "godkend_alle"}, follow=True)
        self.assertContains(svar, "1 medarbejder(e) godkendt.")
        k.refresh_from_db()
        self.assertEqual(k.status, Koersel.Status.GODKENDT)

    def test_kladde_uden_beregning_er_ikke_foraeldet(self):
        k = self.opret_og_upload([("Person 1", "44:12")])
        self.assertFalse(k.beregning_foraeldet)
        self.undtagelser(k, handling="gem")
        k.refresh_from_db()
        self.assertEqual(k.status, Koersel.Status.KLADDE)
        self.assertFalse(k.beregning_foraeldet)

    def test_ingen_skaerm_indeholder_ordet_borger(self):
        def tjek(svar, hvad):
            self.assertEqual(svar.status_code, 200, hvad)
            self.assertNotIn("borger", synlig_tekst(svar.content.decode()).lower(), hvad)

        # Loginsiden ses uden at være logget ind
        self.client.logout()
        svar = self.client.get(reverse("login"))
        tjek(svar, "login")
        self.assertContains(svar, "MedarbejderLøn")
        self.client.force_login(self.bruger)

        # Én uden sats giver en advarsel på undtagelser; én ukendt giver afklaring på gennemse
        Borger.objects.create(navn="Person 4", tj_nr="9004", cpr="000000-0004", afstand_km=Decimal("3.2"), transport="selv", zoner=2)
        tjek(self.client.get(reverse("dashboard")), "forside")
        tjek(self.client.get(reverse("borger_liste")), "medarbejderliste")
        tjek(self.client.get(reverse("borger_liste"), {"q": "findes-ikke"}), "tom medarbejderliste")
        tjek(self.client.get(reverse("borger_ny")), "ny medarbejder")
        tjek(self.client.get(reverse("borger_ret", args=[self.p1.pk])), "ret medarbejder")
        svar = self.client.post(reverse("borger_ny"), {"navn": "Person 20", "afstand_km": "1", "traek_procent": "40"}, follow=True)
        tjek(svar, "meddelelse efter oprettelse")
        self.assertContains(svar, "Person 20 er oprettet.")

        k = self.opret_og_upload([("Person 1", "44:12"), ("Person 4", "10:00"), ("Ukendt", "5:00")])
        tjek(self.client.get(reverse("upload", args=[k.pk])), "upload")
        ukendt = k.linjer.get(fil_navn="Ukendt")
        tjek(self.client.get(reverse("gennemse", args=[k.pk])), "gennemse med uafklaret række")
        tjek(self.client.get(reverse("borger_ny"), {"linje": ukendt.pk, "navn": "Ukendt"}), "ny medarbejder fra række")
        svar = self.client.post(reverse("gennemse", args=[k.pk]), {"handling": "afklar", "linje": ukendt.pk, "borger": "999999"}, follow=True)
        tjek(svar, "afklar med ukendt medarbejder")
        self.assertContains(svar, "Vælg en medarbejder i listen.")
        self.client.post(reverse("gennemse", args=[k.pk]), {"handling": "afklar", "linje": ukendt.pk, "borger": self.p2.pk})
        gem = {"handling": "gem"}
        for l in k.linjer.select_related("borger"):
            gem.update({f"afstand_{l.pk}": "7,5" if l == ukendt else str(l.borger.afstand_km), f"transport_{l.pk}": "ingen" if l == ukendt else l.borger.transport})
        svar = self.client.post(reverse("gennemse", args=[k.pk]), gem, follow=True)
        tjek(svar, "gem på gennemse")
        self.assertContains(svar, "Ændringerne er gemt i medarbejderregistret.")
        tjek(self.client.get(reverse("gennemse", args=[k.pk])), "gennemse")

        svar = self.undtagelser(k)
        tjek(svar, "undtagelser afvist")
        self.assertContains(svar, "Beregningen blev ikke kørt, fordi nogle medarbejdere mangler oplysninger.")
        self.assertContains(svar, "Sats pr. time mangler på medarbejderen.")
        tjek(self.client.get(reverse("undtagelser", args=[k.pk])), "undtagelser")

        p4 = Borger.objects.get(navn="Person 4")
        p4.sats = Decimal("20")
        p4.save()
        svar = self.client.get(self.undtagelser(k)["Location"], follow=True)
        tjek(svar, "godkend efter beregning")
        b = k.beregninger.get()
        resultat = b.resultater.get(navn="Person 1")
        tjek(self.client.get(reverse("grundlag", args=[k.pk, resultat.pk])), "beregningsgrundlag")
        tjek(self.client.get(reverse("godkend", args=[k.pk])), "godkend")
        self.assertNotIn("borger", b.snapshot.lower())

        svar = self.client.post(reverse("godkend", args=[k.pk]), {"handling": "godkend", "resultat": resultat.pk}, follow=True)
        tjek(svar, "godkend en")
        self.assertContains(svar, "1 medarbejder(e) godkendt.")
        svar = self.client.get(reverse("eksport", args=[k.pk]), follow=True)
        tjek(svar, "eksport spærret")
        self.assertContains(svar, "Lønudtrækket kan først eksporteres, når alle medarbejdere er godkendt.")
        tjek(self.client.post(reverse("godkend", args=[k.pk]), {"handling": "godkend_alle"}, follow=True), "godkend alle")
        csv_tekst = self.client.get(reverse("eksport", args=[k.pk])).content.decode("utf-8-sig")
        self.assertNotIn("borger", csv_tekst.lower())

        # Administrationen
        admin = get_user_model().objects.create_superuser("admin", password="x-test-adgangskode-4")
        self.client.force_login(admin)
        tjek(self.client.get(reverse("admin:index")), "administration forside")
        tjek(self.client.get(reverse("admin:app_list", args=["register"])), "administration register")
        tjek(self.client.get(reverse("admin:register_borger_changelist")), "administration medarbejdere")
        tjek(self.client.get(reverse("admin:register_borger_add")), "administration ny medarbejder")
        tjek(self.client.get(reverse("admin:register_borger_change", args=[self.p1.pk])), "administration ret medarbejder")
        self.assertContains(self.client.get(reverse("admin:register_borger_changelist")), "Medarbejdere")

        # Skærmkoden
        with open(settings.BASE_DIR / "static" / "js" / "undtagelser.js", encoding="utf-8") as f:
            self.assertNotIn("borger", f.read().lower())

    def test_auditloggen_viser_objekttypen_medarbejder(self):
        # Gamle poster har klassenavnet Borger som data og en gammel handlingstekst; kun visningen skal vise det nye ord
        AuditLog.objects.create(bruger=self.bruger, handling="Oprettede borger", objekt_type="Borger", objekt_id=str(self.p1.pk), detaljer="Person 1")
        AuditLog.objects.create(bruger=self.bruger, handling="Kørte beregning", objekt_type="Koersel", objekt_id="1", detaljer="Person 2")
        admin = get_user_model().objects.create_superuser("admin", password="x-test-adgangskode-4")
        self.client.force_login(admin)
        url = reverse("admin:audit_auditlog_changelist")

        svar = self.client.get(url)
        self.assertEqual(svar.status_code, 200)
        html = svar.content.decode()
        kolonne = re.findall(r'<td class="field-objekt_type_visning">(.*?)</td>', html)
        self.assertEqual(sorted(kolonne), ["Koersel", "Medarbejder"])
        # Filteret viser det nye ord, men filtrerer stadig på den gemte værdi
        filter_links = dict(re.findall(r'<a href="\?objekt_type=([^"&]*)[^"]*"[^>]*>\s*(.*?)\s*</a>', html))
        self.assertEqual(filter_links, {"Borger": "Medarbejder", "Koersel": "Koersel"})
        self.assertNotIn(">Borger<", html)

        svar = self.client.get(url, {"objekt_type": "Borger"})
        self.assertEqual(svar.context["cl"].result_count, 1)
        self.assertEqual(svar.context["cl"].result_list[0].objekt_type, "Borger")

        # Søgning virker uændret
        svar = self.client.get(url, {"q": "Person 2"})
        self.assertEqual([p.objekt_type for p in svar.context["cl"].result_list], ["Koersel"])

        # De gemte data er uændrede
        self.assertTrue(AuditLog.objects.filter(objekt_type="Borger", handling="Oprettede borger").exists())

    def test_rapporten_bruger_ordet_medarbejder(self):
        k = self.opret_og_upload([("Person 1", "44:12"), ("Person 2", "10:00")])
        self.undtagelser(k)
        b = k.beregninger.get()
        self.assertEqual(b.data["rapport_layout"], LAYOUT_VERSION)
        self.assertEqual(LAYOUT_VERSION, 2)
        # Uden komprimering og ASCII85 kan teksten i PDF-filen læses direkte
        with mock.patch.object(rl_config, "pageCompression", 0), mock.patch.object(rl_config, "useA85", 0):
            pdf = lav_rapport(b.data).decode("latin-1").lower()
        self.assertIn("antal medarbejdere", pdf)
        self.assertIn("hver medarbejder", pdf)
        self.assertNotIn("borger", pdf)

    def test_kraever_login(self):
        self.client.logout()
        for navn, args in (("dashboard", []), ("borger_liste", []), ("godkend", [1]), ("rapport", [1])):
            svar = self.client.get(reverse(navn, args=args))
            self.assertEqual(svar.status_code, 302, navn)
            self.assertIn("/konto/login/", svar["Location"])

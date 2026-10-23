import re
from dataclasses import dataclass
from datetime import datetime, time, timedelta

from openpyxl import load_workbook

TID = re.compile(r"^\s*(\d+):(\d{2})(?::\d{2})?\s*$")


class ImportFejl(Exception):
    pass


@dataclass
class Raekke:
    raekke: int
    navn: str
    minutter: int
    dage: int | None = None


def _til_minutter(v) -> int | None:
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    if isinstance(v, str):
        m = TID.match(v)
        if not m:
            raise ValueError
        return int(m.group(1)) * 60 + int(m.group(2))
    if isinstance(v, timedelta):
        return int(v.total_seconds() // 60)
    if isinstance(v, datetime):
        return int((v - datetime(1899, 12, 30)).total_seconds() // 60)
    if isinstance(v, time):
        return v.hour * 60 + v.minute
    if isinstance(v, (int, float)) and not isinstance(v, bool) and v >= 0:
        return int(round(v * 1440))
    raise ValueError


def _til_dage(v) -> int | None:
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    if isinstance(v, str):
        v = v.strip()
        if not v.isascii() or not v.isdigit():
            raise ValueError
        v = int(v)
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    if isinstance(v, bool) or not isinstance(v, int) or not 1 <= v <= 31:
        raise ValueError
    return v


def laes_timefil(fil) -> list[Raekke]:
    """Læser en Excel-fil fra Punchly med kolonnerne 'Navn' og 'Total timer (timer:minutter)'.

    Kolonnen 'Antal dage' (fremmødedage i måneden) er valgfri."""
    navn_fil = getattr(fil, "name", "")
    if not navn_fil.lower().endswith(".xlsx"):
        raise ImportFejl("Filen skal være en Excel-fil i formatet .xlsx.")
    try:
        wb = load_workbook(fil, read_only=True, data_only=True)
        raekker = list(wb.active.iter_rows(values_only=True))
    except Exception:
        raise ImportFejl("Filen kan ikke læses som en Excel-fil (.xlsx). Den kan være beskadiget.")

    navn_kol = tid_kol = dage_kol = header_raekke = None
    for i, raekke in enumerate(raekker[:50]):
        tekster = [str(c).strip().casefold() if c is not None else "" for c in raekke]
        if "navn" in tekster:
            for j, t in enumerate(tekster):
                if t.startswith("total timer") or t == "timer":
                    navn_kol, tid_kol, header_raekke = tekster.index("navn"), j, i
                    dage_kol = tekster.index("antal dage") if "antal dage" in tekster else None
                    break
        if header_raekke is not None:
            break
    if header_raekke is None:
        raise ImportFejl("Kunne ikke finde kolonnerne 'Navn' og 'Total timer' i filen.")

    fundet: list[Raekke] = []
    problemer: list[str] = []
    for i, raekke in enumerate(raekker[header_raekke + 1 :], start=header_raekke + 2):
        navn = str(raekke[navn_kol]).strip() if navn_kol < len(raekke) and raekke[navn_kol] is not None else ""
        vaerdi = raekke[tid_kol] if tid_kol < len(raekke) else None
        if not navn and (vaerdi is None or str(vaerdi).strip() == ""):
            continue
        if not navn:
            problemer.append(f"Række {i}: navn mangler.")
            continue
        try:
            minutter = _til_minutter(vaerdi)
        except ValueError:
            problemer.append(f"Række {i} ({navn}): '{vaerdi}' er ikke et gyldigt tidsformat (timer:minutter).")
            continue
        if minutter is None:
            problemer.append(f"Række {i} ({navn}): timer mangler.")
            continue
        dage_vaerdi = raekke[dage_kol] if dage_kol is not None and dage_kol < len(raekke) else None
        try:
            dage = _til_dage(dage_vaerdi)
        except ValueError:
            problemer.append(f"Række {i} ({navn}): '{dage_vaerdi}' er ikke et gyldigt antal dage (helt tal fra 1 til 31).")
            continue
        fundet.append(Raekke(i, navn, minutter, dage))

    if problemer:
        mere = f" (og {len(problemer) - 5} flere)" if len(problemer) > 5 else ""
        raise ImportFejl(" ".join(problemer[:5]) + mere)
    if not fundet:
        raise ImportFejl("Filen indeholder ingen rækker med navn og timer.")
    return fundet
